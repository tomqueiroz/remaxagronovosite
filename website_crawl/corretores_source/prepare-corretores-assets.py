from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT.parent / "uploaded_files" / "Lista Atualizada - Corretores Agro.xlsx"
SOURCE_DIR = ROOT / "apps" / "client" / "public" / "images" / "corretores"
OUTPUT_DIR = SOURCE_DIR / "web"
DATA_FILE = ROOT / "apps" / "client" / "src" / "data" / "corretores.ts"

PHOTO_MATCHES = {
    "Ana Dark Frota Pereira": "Ana Dark Pereira.jpg",
    "André Carloto do Nascimento": "André Carloto.jpg",
    "André Luís Bissacot": "André Luís Bissacot.jpg",
    "Carlos Alberto Aarestrup Netto": "Beto Aarestrup.jpg",
    "Bruno Ribeiro Lopes": "Bruno Lopes.jpg",
    "Carolina Magalhães Souza": "Carolina Magalhães.jpg",
    "Delfim da Costa Almeida": "Delfim da Costa Almeida.jpg",
    "Edilberto Stein de Quadros": "Edilberto Quadros.jpg",
    "Edione Neri Ferreira": "Edione Neri Ferreira.jpg",
    "Eduardo Callera Pedrosa": "Eduardo Callera Pedrosa.jpg",
    "Eduardo Farias": "Eduardo Farias.jpg",
    "Efraim Will Bezerra Cavalcante": "Efraim Cavalcante.jpg",
    "Fernando Machado Faria dos Santos": "Fernando Faria.jpg",
    "Flávio José de Sousa Pereira": "Flávio Pereira.jpg",
    "Franco Italo Carvalho Rodrigues": "Franco Italo Rodrigues.jpg",
    "Guilherme Guimarães Vilela": "Guilherme Vilela.jpg",
    "Jones Henrique Canova": "Jones Henrique Canova.jpg",
    "José Carlos do Prado Júnior": "José Carlos do Prado.jpg",
}

STATE_NAMES = {
    "AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas",
    "BA": "Bahia", "CE": "Ceará", "DF": "Distrito Federal",
    "ES": "Espírito Santo", "GO": "Goiás", "MA": "Maranhão",
    "MT": "Mato Grosso", "MS": "Mato Grosso do Sul", "MG": "Minas Gerais",
    "PA": "Pará", "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco",
    "PI": "Piauí", "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte",
    "RS": "Rio Grande do Sul", "RO": "Rondônia", "RR": "Roraima",
    "SC": "Santa Catarina", "SP": "São Paulo", "SE": "Sergipe",
    "TO": "Tocantins",
}


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    return f"+55{digits}" if digits else ""


def load_records() -> list[dict[str, str]]:
    workbook = load_workbook(XLSX, read_only=True, data_only=False)
    worksheet = workbook["Sheet1"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(value).strip() for value in rows[0]]
    records: list[dict[str, str]] = []
    for index, row in enumerate(rows[1:], start=1):
        source = {header: str(value).strip() if value is not None else "" for header, value in zip(headers, row)}
        name = source["AGENTE"]
        photo_source = PHOTO_MATCHES.get(name)
        photo = f"/images/corretores/web/{slugify(name)}.webp" if photo_source else None
        records.append({
            "id": f"corretor-{index:02d}",
            "state": source["ESTADO"],
            "name": name,
            "phone": source["TELEFONE"],
            "phoneHref": normalize_phone(source["TELEFONE"]),
            "email": source["EMAIL"],
            "city": source["CIDADE"],
            "photo": photo,
        })
    return records


def optimize_photos() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    expected = set()
    for name, source_name in PHOTO_MATCHES.items():
        source = SOURCE_DIR / source_name
        destination = OUTPUT_DIR / f"{slugify(name)}.webp"
        expected.add(destination.name)
        with Image.open(source) as image:
            image = image.convert("RGB")
            image.thumbnail((900, 1350), Image.Resampling.LANCZOS)
            image.save(destination, "WEBP", quality=82, method=6)
    for path in OUTPUT_DIR.glob("*.webp"):
        if path.name not in expected:
            path.unlink()


def write_placeholder() -> None:
    placeholder = OUTPUT_DIR / "corretor-placeholder.svg"
    placeholder.write_text(
        """<svg xmlns=\"http://www.w3.org/2000/svg\" viewBox=\"0 0 600 800\" role=\"img\" aria-labelledby=\"title desc\">\n"
        "  <title id=\"title\">Retrato genérico de corretor REMAX Agro</title>\n"
        "  <desc id=\"desc\">Ilustração neutra usada quando não há fotografia disponível.</desc>\n"
        "  <rect width=\"600\" height=\"800\" fill=\"#eef1f4\"/>\n"
        "  <circle cx=\"300\" cy=\"280\" r=\"112\" fill=\"#c5cdd6\"/>\n"
        "  <path d=\"M112 690c20-143 98-218 188-218s168 75 188 218\" fill=\"#c5cdd6\"/>\n"
        "  <path d=\"M170 690h260\" stroke=\"#0c2749\" stroke-width=\"10\" stroke-linecap=\"round\" opacity=\".22\"/>\n"
        "  <text x=\"300\" y=\"746\" text-anchor=\"middle\" font-family=\"Roboto,Arial,sans-serif\" font-size=\"24\" font-weight=\"700\" letter-spacing=\"2\" fill=\"#0c2749\">REMAX AGRO</text>\n"
        "</svg>\n""",
        encoding="utf-8",
    )


def write_dataset(records: list[dict[str, str]]) -> None:
    states = [{"uf": uf, "name": name} for uf, name in STATE_NAMES.items()]
    data_json = json.dumps(records, ensure_ascii=False, indent=2)
    states_json = json.dumps(states, ensure_ascii=False, indent=2)
    content = f"""/* @section: corretores-dataset */
export type Corretor = {{
  id: string
  state: string
  name: string
  phone: string
  phoneHref: string
  email: string
  city: string
  photo: string | null
}}

export type BrazilState = {{
  uf: string
  name: string
}}

export const BRAZIL_STATES: BrazilState[] = {states_json}

export const CORRETORES: Corretor[] = {data_json}

export const CORRETORES_BY_STATE = CORRETORES.reduce<Record<string, Corretor[]>>((groups, corretor) => {{
  const current = groups[corretor.state] ?? []
  current.push(corretor)
  groups[corretor.state] = current
  return groups
}}, {{}})

export const COVERED_STATES = new Set(Object.keys(CORRETORES_BY_STATE))
export const CORRETOR_PLACEHOLDER = '/images/corretores/web/corretor-placeholder.svg'
"""
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    DATA_FILE.write_text(content, encoding="utf-8")


def main() -> None:
    records = load_records()
    optimize_photos()
    write_placeholder()
    write_dataset(records)
    counts: dict[str, int] = {}
    for record in records:
        counts[record["state"]] = counts.get(record["state"], 0) + 1
    photo_count = sum(1 for record in records if record["photo"])
    total_bytes = sum(path.stat().st_size for path in OUTPUT_DIR.glob("*.webp"))
    print(json.dumps({
        "records": len(records),
        "states": len(counts),
        "counts": counts,
        "photos": photo_count,
        "placeholders": len(records) - photo_count,
        "optimized_total_bytes": total_bytes,
        "data_file": str(DATA_FILE.relative_to(ROOT)),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
